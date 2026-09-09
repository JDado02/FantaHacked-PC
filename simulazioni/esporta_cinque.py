# -*- coding: utf-8 -*-
"""Le cinque aste giocate dal posto di Davide, in un foglio ciascuna.

Il primo foglio e' la rosa che il motore comprerebbe se ogni giocatore
costasse quello che vale: e' il piano, non un risultato. Gli altri cinque sono
quello che il piano diventa quando in mezzo ci sono sette avversari che
rilanciano.
"""
import os, shutil, sys

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, 'motore'))

import percorsi, db as dbmod, regole as regmod
from asta import StatoAsta
from valutazione import Valutatore
from ottimizzatore import Ottimizzatore
from cinque_aste import gioca, undici_atteso, NOMI, RUOLI

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NOME_RUOLO = {'P': 'Portieri', 'D': 'Difensori',
              'C': 'Centrocampisti', 'A': 'Attaccanti'}
SCURO = PatternFill('solid', fgColor='1F2A36')
BANDA = PatternFill('solid', fgColor='F2F5F8')
REPARTO = PatternFill('solid', fgColor='E4EAF0')
MIO = PatternFill('solid', fgColor='FFF4D6')
BIANCO = Font(color='FFFFFF', bold=True, size=11)
GRASSETTO = Font(bold=True)
ORO = Font(bold=True, color='9A6B00')
SOTTILE = Side(style='thin', color='D3DAE1')
BORDO = Border(left=SOTTILE, right=SOTTILE, top=SOTTILE, bottom=SOTTILE)


def intesta(ws, riga, colonne):
    for c, testo in enumerate(colonne, start=1):
        cella = ws.cell(row=riga, column=c, value=testo)
        cella.fill = SCURO
        cella.font = BIANCO
        cella.alignment = Alignment(horizontal='center', vertical='center')
        cella.border = BORDO
    ws.row_dimensions[riga].height = 20
    return riga + 1


def larghezze(ws, misure):
    for i, w in enumerate(misure, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def rosa_ideale():
    dst = os.path.join(QUI, 'ideale.db')
    shutil.copyfile(percorsi.DB_FILE, dst)
    con = dbmod.connetti(dst)
    reg = regmod.carica()
    st = StatoAsta(con, reg)
    st.inizializza(NOMI[1:], mio_nome=NOMI[0])
    v = Valutatore(con, reg, st)
    o = Ottimizzatore(v)
    piano = o.piano()
    fuori = []
    for r in RUOLI:
        d = piano['per_ruolo'][r]
        righe = []
        for t in d['obiettivi']:
            x = v.g[t['id']]
            righe.append({'nome': x.nome, 'squadra': x.squadra,
                          'prezzo': t['costo'], 'fm': round(x.fm, 2),
                          'presenze': int(round(x.presenze)),
                          'punti': round(x.presenze * x.fm),
                          'grado': x.grado or '?'})
        righe.sort(key=lambda z: -z['prezzo'])
        fuori.append((r, d['crediti'], righe))
    con.close()
    try:
        os.remove(dst)
    except OSError:
        pass
    return fuori


def foglio_ideale(wb, dati):
    ws = wb.create_sheet('Rosa ideale')
    ws['A1'] = 'La rosa che il motore comprerebbe ai suoi prezzi'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = ("E' il piano a inizio asta, non un risultato: presuppone di "
                "portare via ognuno di questi al prezzo indicato. In asta vera "
                "non succede quasi mai, ed e' proprio per questo che il motore "
                "ricalcola tutto dopo ogni acquisto.")
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A2:G2')
    ws.row_dimensions[2].height = 32
    larghezze(ws, [26, 15, 9, 8, 8, 9, 15])
    riga = 4
    totale = 0
    for ruolo, crediti, righe in dati:
        cella = ws.cell(row=riga, column=1,
                        value='%s — %d crediti' % (NOME_RUOLO[ruolo], crediti))
        cella.font = GRASSETTO
        for c in range(1, 8):
            ws.cell(row=riga, column=c).fill = REPARTO
        riga += 1
        riga = intesta(ws, riga, ['Giocatore', 'Squadra', 'Prezzo', 'Fm',
                                  'Presenze', 'Punti', 'Ruolo in squadra'])
        for i, d in enumerate(righe):
            valori = [d['nome'], d['squadra'], d['prezzo'], d['fm'],
                      d['presenze'], d['punti'], d['grado']]
            for c, val in enumerate(valori, start=1):
                cella = ws.cell(row=riga, column=c, value=val)
                cella.border = BORDO
                if i % 2:
                    cella.fill = BANDA
                if c == 3:
                    cella.font = ORO
            riga += 1
            totale += d['prezzo']
        riga += 1
    ws.cell(row=riga, column=1, value='Totale speso').font = GRASSETTO
    ws.cell(row=riga, column=3, value=totale).font = ORO
    ws.freeze_panes = 'A4'


def foglio_asta(wb, n, rose):
    ws = wb.create_sheet('Asta %d' % n)
    mia = rose[NOMI[0]]
    spesa = dict((r, sum(d['prezzo'] for d in mia[r])) for r in RUOLI)
    ws['A1'] = 'Asta %d — la rosa che è uscita' % n
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = ('Spesa per reparto: portieri %d, difensori %d, centrocampisti '
                '%d, attaccanti %d. Totale %d su 500.'
                % (spesa['P'], spesa['D'], spesa['C'], spesa['A'],
                   sum(spesa.values())))
    larghezze(ws, [26, 15, 9, 9, 8, 8, 9, 15])
    riga = 4
    for ruolo in RUOLI:
        cella = ws.cell(row=riga, column=1,
                        value='%s — %d crediti' % (NOME_RUOLO[ruolo], spesa[ruolo]))
        cella.font = GRASSETTO
        for c in range(1, 9):
            ws.cell(row=riga, column=c).fill = REPARTO
        riga += 1
        riga = intesta(ws, riga, ['Giocatore', 'Squadra', 'Pagato', 'Valeva',
                                  'Fm', 'Presenze', 'Punti', 'Ruolo in squadra'])
        for i, d in enumerate(mia[ruolo]):
            valori = [d['nome'], d['squadra'], d['prezzo'], d['base'], d['fm'],
                      d['presenze'], d['punti'], d['grado']]
            for c, val in enumerate(valori, start=1):
                cella = ws.cell(row=riga, column=c, value=val)
                cella.border = BORDO
                if i % 2:
                    cella.fill = BANDA
                if c == 3:
                    cella.font = ORO
            riga += 1
        riga += 1

    riga += 1
    ws.cell(row=riga, column=1,
            value='Come finisce la lega (punti dell\'undici titolare atteso)'
            ).font = GRASSETTO
    riga += 1
    riga = intesta(ws, riga, ['Pos', 'Squadra', 'Punti undici', 'Spesa'])
    classifica = sorted(((undici_atteso(rose[n_]), n_) for n_ in NOMI),
                        reverse=True)
    for posto, (punti, nome) in enumerate(classifica, start=1):
        speso = sum(d['prezzo'] for r in RUOLI for d in rose[nome][r])
        for c, val in enumerate([posto, nome, round(punti), speso], start=1):
            cella = ws.cell(row=riga, column=c, value=val)
            cella.border = BORDO
            if nome == NOMI[0]:
                cella.fill = MIO
                cella.font = GRASSETTO
        riga += 1
    ws.freeze_panes = 'A4'


def main():
    wb = Workbook()
    wb.remove(wb.active)
    print('rosa ideale...')
    foglio_ideale(wb, rosa_ideale())
    for i, seme in enumerate([101, 202, 303, 404, 505], start=1):
        print('asta %d...' % i)
        foglio_asta(wb, i, gioca(seme))
    fuori = os.path.join(QUI, 'Cinque_aste_dal_tuo_posto.xlsx')
    wb.save(fuori)
    print('scritto', fuori)


if __name__ == '__main__':
    main()
